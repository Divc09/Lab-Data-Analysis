from __future__ import annotations

import csv
import json
import math
import shutil
import textwrap
from collections import Counter, defaultdict
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ANALYSIS = Path(r"D:\BacklundLabResearch\Data\Experiments\429_TDD\analysis_outputs\429_TDD_full_characterization")
OUT_DIR = Path(r"D:\BacklundLabResearch\Data\Experiments\429_TDD\analysis_outputs\429_TDD_research_deck")
PPTX_PATH = OUT_DIR / "429_TDD_NV_sensing_research_deck.pptx"
PREVIEW_DIR = OUT_DIR / "preview_png"

WIDE_W, WIDE_H = 13.333333, 7.5
BG = "F7F7F2"
INK = "172126"
MUTED = "687274"
ACCENT = "167C80"
ACCENT2 = "B85C38"
BLUE = "2D5B88"
RULE = "CBD1CA"
SOFT = "E9ECE4"


def rgb(hex_color: str) -> RGBColor:
    h = hex_color.strip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def fnum(value: object, digits: int = 3) -> str:
    try:
        v = float(value)
    except Exception:
        return ""
    if not math.isfinite(v):
        return ""
    if abs(v) >= 1000:
        return f"{v:,.0f}"
    if abs(v) < 0.01 and v != 0:
        return f"{v:.2e}"
    return f"{v:.{digits}g}"


def load_csv(name: str) -> list[dict[str, str]]:
    with (ANALYSIS / "tables" / name).open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def load_data() -> dict[str, object]:
    summary = json.loads((ANALYSIS / "tables" / "summary.json").read_text(encoding="utf-8"))
    best = load_csv("best_metrics.csv")
    fit = load_csv("fit_summary.csv")
    inventory = load_csv("inventory.csv")
    iteration = load_csv("iteration_summary.csv")
    noise = load_csv("noise_spectrum_summary.csv")
    return {
        "summary": summary,
        "best": best,
        "fit": fit,
        "inventory": inventory,
        "iteration": iteration,
        "noise": noise,
    }


def set_bg(slide, color: str = BG) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def box(slide, x, y, w, h, color=SOFT, line=None, radius=False):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(0.8)
    else:
        shape.line.fill.background()
    return shape


def add_text(slide, text, x, y, w, h, size=24, color=INK, bold=False, align=None, font="Aptos"):
    tx = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tx.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    if align:
        p.alignment = align
    run = p.add_run()
    run.text = str(text)
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return tx


def add_bullets(slide, bullets, x, y, w, h, size=18, color=INK, gap=4):
    tx = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tx.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = str(item)
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = rgb(color)
        p.space_after = Pt(gap)
        p.level = 0
    return tx


def title(slide, section: str, heading: str, sub: str = "") -> None:
    add_text(slide, section.upper(), 0.65, 0.28, 2.5, 0.25, 9, ACCENT, True)
    add_text(slide, heading, 0.65, 0.58, 8.8, 0.58, 26, INK, True)
    if sub:
        add_text(slide, sub, 0.67, 1.13, 9.6, 0.38, 12.5, MUTED)
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.65), Inches(1.56), Inches(12.0), Inches(0.015))
    line.fill.solid()
    line.fill.fore_color.rgb = rgb(RULE)
    line.line.fill.background()


def footer(slide, note: str = "Source: 429_TDD .mat analysis outputs generated from nvfit pipeline.") -> None:
    add_text(slide, note, 0.65, 7.13, 8.0, 0.2, 8.2, MUTED)


def add_metric(slide, label, value, x, y, w, h=0.72, color=ACCENT):
    add_text(slide, str(value), x, y, w, h * 0.6, 24, color, True)
    add_text(slide, label, x, y + h * 0.48, w, h * 0.45, 8.5, MUTED)


def image_fit(slide, path: Path, x, y, w, h):
    if not path.exists():
        box(slide, x, y, w, h, "F0D7D2", "D58A7B")
        add_text(slide, f"Missing figure\n{path.name}", x + 0.2, y + 0.2, w - 0.4, h - 0.4, 12, "7A2E25")
        return None
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def add_bar_chart(slide, title_text, categories, values, x, y, w, h, color=ACCENT):
    chart_data = CategoryChartData()
    chart_data.categories = categories
    chart_data.add_series(title_text, values)
    frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), chart_data)
    chart = frame.chart
    chart.has_legend = False
    chart.chart_title.has_text_frame = True
    chart.chart_title.text_frame.text = title_text
    chart.value_axis.has_major_gridlines = True
    chart.category_axis.tick_labels.font.size = Pt(8)
    chart.value_axis.tick_labels.font.size = Pt(8)
    try:
        chart.plots[0].series[0].format.fill.solid()
        chart.plots[0].series[0].format.fill.fore_color.rgb = rgb(color)
    except Exception:
        pass
    return frame


def add_line_chart(slide, title_text, categories, series_map, x, y, w, h):
    chart_data = CategoryChartData()
    chart_data.categories = categories
    for name, values in series_map.items():
        chart_data.add_series(str(name), values)
    frame = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(x), Inches(y), Inches(w), Inches(h), chart_data)
    chart = frame.chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.chart_title.has_text_frame = True
    chart.chart_title.text_frame.text = title_text
    chart.category_axis.tick_labels.font.size = Pt(8)
    chart.value_axis.tick_labels.font.size = Pt(8)
    return frame


def add_table(slide, headers, rows, x, y, w, h, font_size=8.8):
    table_shape = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h))
    table = table_shape.table
    for idx, header in enumerate(headers):
        cell = table.cell(0, idx)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(INK)
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(font_size)
                r.font.bold = True
                r.font.color.rgb = rgb("FFFFFF")
    for r_idx, row in enumerate(rows, 1):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = str(value)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb("FFFFFF" if r_idx % 2 else "F0F2EC")
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(font_size)
                    run.font.color.rgb = rgb(INK)
    return table_shape


def fig(name: str) -> Path:
    return ANALYSIS / "figures" / "curated" / name


def per_file_path(relative_path: str, figure_by_rel: dict[str, str]) -> Path:
    rel = figure_by_rel.get(relative_path, "")
    return ANALYSIS / rel.replace("/", "\\")


def rows_for(best: list[dict[str, str]], exp: str, limit=5) -> list[dict[str, str]]:
    return [r for r in best if r["experiment_type"] == exp][:limit]


def make_preview_contact(slide_titles: list[str], image_paths: list[Path]) -> Path:
    PREVIEW_DIR.mkdir(parents=True, exist_ok=True)
    thumbs = []
    font = ImageFont.load_default()
    for i, name in enumerate(slide_titles, 1):
        im = Image.new("RGB", (480, 270), "#" + BG)
        d = ImageDraw.Draw(im)
        d.rectangle([0, 0, 479, 269], outline="#" + RULE)
        d.text((18, 18), f"{i:02d}", fill="#" + ACCENT, font=font)
        wrapped = textwrap.wrap(name, width=42)[:3]
        d.text((56, 18), "\n".join(wrapped), fill="#" + INK, font=font)
        if i - 1 < len(image_paths) and image_paths[i - 1].exists():
            try:
                pic = Image.open(image_paths[i - 1]).convert("RGB")
                pic.thumbnail((410, 160))
                im.paste(pic, (35, 92))
            except Exception:
                pass
        thumbs.append(im)
    cols = 3
    rows = math.ceil(len(thumbs) / cols)
    sheet = Image.new("RGB", (cols * 480, rows * 270), "white")
    for idx, im in enumerate(thumbs):
        sheet.paste(im, ((idx % cols) * 480, (idx // cols) * 270))
    out = PREVIEW_DIR / "deck_contact_sheet.png"
    sheet.save(out)
    return out


def build_deck() -> tuple[list[str], list[Path]]:
    data = load_data()
    summary = data["summary"]
    best = data["best"]
    fit = data["fit"]
    inventory = data["inventory"]
    iteration = data["iteration"]
    noise = data["noise"]

    prs = Presentation()
    prs.slide_width = Inches(WIDE_W)
    prs.slide_height = Inches(WIDE_H)
    blank = prs.slide_layouts[6]
    slide_titles: list[str] = []
    preview_imgs: list[Path] = []

    def new_slide(name: str):
        slide = prs.slides.add_slide(blank)
        set_bg(slide)
        slide_titles.append(name)
        preview_imgs.append(Path(""))
        return slide

    def preview(path: Path) -> None:
        if preview_imgs:
            preview_imgs[-1] = path

    by_type = summary["by_type"]
    figure_by_rel = {r["relative_path"]: r.get("figure", "") for r in fit}
    counts = {k: v["n"] for k, v in by_type.items()}
    med_r2 = {k: v["median_r2"] for k, v in by_type.items() if v["median_r2"] is not None}

    # 1 Cover
    slide = new_slide("429_TDD NV Sensing Diamond Characterization")
    box(slide, 0, 0, 13.333, 7.5, BG)
    add_text(slide, "429_TDD", 0.75, 0.62, 3.1, 0.5, 18, ACCENT, True)
    add_text(slide, "NV sensing diamond\nresearch characterization", 0.75, 1.18, 7.6, 1.6, 40, INK, True)
    add_text(slide, "Rabi, Ramsey, spin echo, T1, XY8, DEER, ODMR, and spatial context from 255 MATLAB runs", 0.78, 3.05, 7.5, 0.62, 17, MUTED)
    image_fit(slide, fig("nv_metric_distributions.png"), 7.95, 0.75, 4.7, 3.2)
    image_fit(slide, fig("Ramsey_top_overlay.png"), 7.95, 4.1, 4.7, 2.55)
    preview(fig("Ramsey_top_overlay.png"))
    add_text(slide, "Generated from analyzed .mat data\nBacklund Lab analysis workspace", 0.78, 6.35, 4.4, 0.45, 11, MUTED)

    # 2 Executive signal
    slide = new_slide("Executive read: the diamond is measurable, but contrast and sequence quality vary strongly")
    title(slide, "summary", "The dataset supports a full sensing characterization, with uneven sequence quality")
    add_metric(slide, "MAT files loaded", summary["total_files"], 0.8, 1.9, 1.5)
    add_metric(slide, "Successful fits", summary["fit_ok"], 2.45, 1.9, 1.5)
    add_metric(slide, "Context / plot-only", summary["plot_only"], 4.1, 1.9, 1.7)
    add_metric(slide, "Single-point XY8", summary["skipped"], 5.95, 1.9, 1.6)
    add_bullets(
        slide,
        [
            "Best calibration-quality Rabi gives pi ~24.6 ns after combiner; earlier runs span ~6-18 ns, consistent with changing drive conditions.",
            "Best Ramsey shows T2* ~174 ns at ~11.6 MHz detuning; many Ramsey runs are lower-R2, so phase/reference details matter.",
            "Spin echo clusters near T2 ~0.8-1.1 us; XY8 decay fits extend to roughly 1.2-3.6 us but with lower confidence at high pulse count.",
            "DEER frequency centers around ~0.741 GHz in the cleaner sweeps; duration and frequency sweeps are useful but not uniformly high-contrast.",
        ],
        0.85,
        3.25,
        11.5,
        2.3,
        16,
    )
    footer(slide)

    # 3 Dataset composition
    slide = new_slide("Dataset composition and analysis triage")
    title(slide, "scope", "255 files were sorted into fit-ready sensing data and context scans")
    cats = ["ODMR", "Line", "XY8/DD", "Rabi", "Ramsey", "DEER F", "DEER Dur", "Echo", "T1", "DEER Pos"]
    vals = [
        counts.get("ODMR", 0),
        counts.get("LineScan", 0),
        counts.get("DynamicDecoupling", 0),
        counts.get("Rabi", 0),
        counts.get("Ramsey", 0),
        counts.get("DEERFrequency", 0),
        counts.get("DEERDuration", 0),
        counts.get("SpinEcho", 0),
        counts.get("T1", 0),
        counts.get("DEERPosition", 0),
    ]
    add_bar_chart(slide, "File count by experiment family", cats, vals, 0.75, 1.9, 7.4, 4.35, ACCENT)
    preview(fig("nv_metric_distributions.png"))
    add_bullets(
        slide,
        [
            f"{summary['alignment_like']} files flagged as alignment-like context; most are ODMR/magnet-positioning repeats.",
            f"{summary['checkpoints']} checkpoint files preserved in the inventory; final runs are preferred when duplicates exist.",
            "Sensing interpretation is weighted toward sequence families with real sweeps and repeatability information.",
        ],
        8.55,
        2.05,
        3.8,
        2.8,
        16,
    )
    footer(slide)

    # 4 Workflow
    slide = new_slide("Analysis workflow")
    title(slide, "method", "The report separates physics extraction from acquisition context")
    steps = [
        ("Load", "All MATLAB files loaded through nvfit snapshot/schema handling."),
        ("Classify", "Experiment families inferred from scan metadata, notes, and filename."),
        ("Fit", "Sequence-specific models used for Rabi, Ramsey, echo/DD, T1, ODMR, and DEER."),
        ("Curate", "Alignment ODMR summarized lightly; sensing sequences compared across runs and iterations."),
        ("Report", "Per-file plots, overlays, fit tables, iteration stability, FFT diagnostics, and final synthesis."),
    ]
    x = 0.8
    for idx, (head, body) in enumerate(steps):
        add_text(slide, str(idx + 1), x, 2.0, 0.38, 0.32, 16, "FFFFFF", True, PP_ALIGN.CENTER)
        box(slide, x - 0.06, 2.38, 2.16, 2.05, "FFFFFF", RULE, True)
        add_text(slide, head, x + 0.12, 2.56, 1.8, 0.28, 18, ACCENT, True)
        add_text(slide, body, x + 0.12, 3.02, 1.75, 0.95, 10.5, INK)
        x += 2.45
    footer(slide)

    # 5 Fit quality
    slide = new_slide("Fit quality map")
    title(slide, "quality", "High R2 does not always mean high scientific value, but it flags reliable extraction")
    qcats = list(med_r2.keys())
    qvals = [float(med_r2[k]) for k in qcats]
    add_bar_chart(slide, "Median R2 by fit family", qcats, qvals, 0.72, 1.9, 7.6, 4.2, BLUE)
    add_bullets(
        slide,
        [
            "ODMR, DEER decay, spin echo, and Rabi contain the most stable fit surfaces.",
            "Ramsey and DEER frequency/duration contain useful structure but require visual inspection of traces and residuals.",
            "T1 fits are split: normalized/balanced runs fit well, while raw long scans are poorly described by the selected model.",
        ],
        8.65,
        2.0,
        3.65,
        2.6,
        15.5,
    )
    footer(slide)

    # 6 ODMR
    slide = new_slide("ODMR establishes resonance context, not the central result")
    title(slide, "odmr", "Representative ODMR is strong; repeated alignment ODMR was intentionally downweighted")
    image_fit(slide, fig("ODMR_top_overlay.png"), 0.75, 1.85, 6.2, 3.85)
    image_fit(slide, fig("odmr_center_timeline.png"), 7.25, 1.85, 5.3, 2.15)
    preview(fig("ODMR_top_overlay.png"))
    odmr = rows_for(best, "ODMR", 3)
    add_table(
        slide,
        ["Run", "Center", "FWHM", "R2"],
        [[Path(r["relative_path"]).stem[:32], fnum(r["center_freq_units"]), fnum(r["fwhm_units"]), fnum(r["r2"])] for r in odmr],
        7.25,
        4.28,
        5.25,
        1.2,
        7.8,
    )
    add_text(slide, "Interpretation: resonance identification is robust enough for downstream sequence tuning, but the many magnet-alignment files should not be counted as independent material characterization.", 0.82, 6.08, 11.5, 0.55, 13, INK)
    footer(slide)

    # 7 Rabi
    slide = new_slide("Rabi calibration")
    title(slide, "rabi", "Rabi runs show clear drive calibration, with drive-condition dependence")
    image_fit(slide, fig("Rabi_top_overlay.png"), 0.72, 1.85, 6.45, 4.25)
    preview(fig("Rabi_top_overlay.png"))
    rabi = rows_for(best, "Rabi", 5)
    add_bar_chart(slide, "Top Rabi pi time [ns]", [str(i) for i in range(1, len(rabi) + 1)], [float(r["pi_time_ns"]) for r in rabi], 7.55, 2.05, 4.75, 2.55, ACCENT)
    add_table(
        slide,
        ["Rank", "Model", "Pi ns", "R2"],
        [[r["rank"], r["model"].replace("Rabi", ""), fnum(r["pi_time_ns"]), fnum(r["r2"])] for r in rabi],
        7.55,
        4.86,
        4.75,
        1.15,
        8.4,
    )
    footer(slide)

    # 8 Rabi science read
    slide = new_slide("Rabi interpretation")
    title(slide, "rabi", "The best Rabi run is not just the strongest oscillation; it is the cleanest calibration state")
    add_bullets(
        slide,
        [
            "The best post-combiner/ref-peak run gives pi ~24.6 ns at R2 ~0.987, a credible working calibration for subsequent DEER and echo sequences.",
            "Shorter apparent pi times in earlier May 5-7 runs indicate substantially different microwave delivery or power conditions.",
            "AOM compensation and repolarization variations change both contrast baseline and oscillation damping; compare these before selecting a final pulse length.",
            "Use the best-fit table for calibration candidates, but inspect per-file plots for beating or chirp-like behavior before hard-locking pulse timing.",
        ],
        0.85,
        1.95,
        5.5,
        3.9,
        17,
    )
    rabi_detail = per_file_path(rabi[0]["relative_path"], figure_by_rel)
    image_fit(slide, rabi_detail, 6.85, 1.85, 5.6, 4.25)
    preview(rabi_detail)
    footer(slide)

    # 9 Ramsey
    slide = new_slide("Ramsey dephasing and detuning")
    title(slide, "ramsey", "Ramsey shows short T2* and strong sensitivity to phase/reference choices")
    image_fit(slide, fig("Ramsey_top_overlay.png"), 0.72, 1.85, 6.55, 4.35)
    preview(fig("Ramsey_top_overlay.png"))
    ramsey = rows_for(best, "Ramsey", 5)
    add_bar_chart(slide, "Top Ramsey T2* [ns]", [str(i) for i in range(1, len(ramsey) + 1)], [float(r["T2_star_ns"]) for r in ramsey], 7.62, 2.0, 4.7, 2.45, BLUE)
    add_table(
        slide,
        ["Run", "T2* ns", "det MHz", "R2"],
        [[r["rank"], fnum(r["T2_star_ns"]), fnum(r["detuning_MHz"]), fnum(r["r2"])] for r in ramsey],
        7.62,
        4.82,
        4.7,
        1.25,
        8.4,
    )
    footer(slide)

    # 10 Ramsey interpretation
    slide = new_slide("Ramsey interpretation")
    title(slide, "ramsey", "The cleanest Ramsey trace supports T2* ~174 ns; lower-quality traces are still diagnostic")
    ramsey_detail = per_file_path(ramsey[0]["relative_path"], figure_by_rel)
    image_fit(slide, ramsey_detail, 0.72, 1.82, 6.2, 4.35)
    preview(ramsey_detail)
    add_bullets(
        slide,
        [
            "Best file: May9 10 MHz detuned Ramsey, hyperfine model, R2 ~0.87, T2* ~174 ns, detuning ~11.6 MHz.",
            "May10 +X/-X files provide phase-control checks at RF ~2.140 GHz but fit lower than the best May9 run.",
            "Median Ramsey R2 is low compared with Rabi/echo, pointing to real sensitivity to pulse phasing, reference construction, and short dephasing window.",
            "Scientifically, Ramsey constrains near-term DC/slow-noise sensing expectations: usable but short-lived free precession contrast.",
        ],
        7.25,
        1.95,
        4.95,
        3.6,
        16,
    )
    footer(slide)

    # 11 Spin echo
    slide = new_slide("Spin echo coherence")
    title(slide, "spin echo", "Echo extends usable coherence to roughly the microsecond scale")
    image_fit(slide, fig("SpinEcho_top_overlay.png"), 0.72, 1.85, 6.35, 4.25)
    preview(fig("SpinEcho_top_overlay.png"))
    echo = rows_for(best, "SpinEcho", 5)
    add_bar_chart(slide, "Top spin echo T2 [ns]", [str(i) for i in range(1, len(echo) + 1)], [float(r["T2_ns"]) for r in echo], 7.48, 2.0, 4.8, 2.55, ACCENT)
    add_bullets(
        slide,
        [
            "Top echo fits cluster from ~0.78 to 1.13 us.",
            "The best file is a checkpoint, so inspect its completion metadata before treating it as final.",
            "Echo performance is substantially longer than Ramsey T2*, consistent with refocusing of quasi-static dephasing.",
        ],
        7.55,
        4.9,
        4.6,
        1.3,
        13.5,
    )
    footer(slide)

    # 12 XY8/DD
    slide = new_slide("Dynamic decoupling")
    title(slide, "xy8", "XY8 improves the apparent coherence envelope but higher-N confidence is mixed")
    image_fit(slide, fig("DynamicDecoupling_top_overlay.png"), 0.72, 1.85, 6.45, 4.25)
    preview(fig("DynamicDecoupling_top_overlay.png"))
    dd = rows_for(best, "DynamicDecoupling", 5)
    add_bar_chart(slide, "Top XY8/DD T2 [ns]", [str(i) for i in range(1, len(dd) + 1)], [float(r["T2_dd_ns"]) for r in dd], 7.55, 2.0, 4.75, 2.55, BLUE)
    add_table(
        slide,
        ["Rank", "T2 ns", "R2"],
        [[r["rank"], fnum(r["T2_dd_ns"]), fnum(r["r2"])] for r in dd],
        7.55,
        4.9,
        4.0,
        1.08,
        8.8,
    )
    footer(slide)

    # 13 Noise spectrum
    slide = new_slide("XY8 noise-spectrum points")
    title(slide, "xy8", "Single-point XY8 scans map contrast response versus filter frequency")
    image_fit(slide, fig("noise_spectrum_xy8_summary.png"), 0.72, 1.85, 6.1, 4.25)
    preview(fig("noise_spectrum_xy8_summary.png"))
    grouped = defaultdict(dict)
    for r in noise:
        grouped[int(float(r["xy8_repetitions"]))][float(r["filter_frequency_mhz"])] = float(r["mean_contrast"])
    freqs = sorted({float(r["filter_frequency_mhz"]) for r in noise})
    series = {f"XY8-{n}": [grouped[n].get(f, None) for f in freqs] for n in sorted(grouped)}
    add_line_chart(slide, "Native chart: mean contrast by filter frequency", [str(f) for f in freqs], series, 7.1, 2.0, 5.1, 3.3)
    add_text(slide, "These are single-point response measurements, not decay curves; the correct interpretation is contrast response versus filter function, not stretched-exponential T2.", 7.2, 5.55, 4.8, 0.55, 12.5, INK)
    footer(slide)

    # 14 T1
    slide = new_slide("T1 relaxation")
    title(slide, "t1", "T1 results are highly model-dependent; normalized/balanced traces are most credible")
    image_fit(slide, fig("T1_top_overlay.png"), 0.72, 1.85, 6.4, 4.25)
    preview(fig("T1_top_overlay.png"))
    t1 = rows_for(best, "T1", 5)
    add_bar_chart(slide, "Top T1 estimates [us]", [str(i) for i in range(1, len(t1) + 1)], [float(r["T1_ns"]) / 1000.0 for r in t1], 7.55, 2.0, 4.75, 2.55, ACCENT2)
    add_bullets(
        slide,
        [
            "Best normalized/balanced T1: ~254 us with R2 ~0.83.",
            "Some raw long scans yield multi-ms estimates with near-zero or negative R2; those should be treated as failed model descriptions, not physical T1 claims.",
            "Recommendation: repeat T1 with balanced/normalized acquisition and explicit dark-time axis verification.",
        ],
        7.55,
        4.88,
        4.6,
        1.35,
        13.3,
    )
    footer(slide)

    # 15 DEER frequency
    slide = new_slide("DEER frequency sweeps")
    title(slide, "deer", "DEER frequency sweeps locate a repeatable response near 0.741 GHz")
    image_fit(slide, fig("DEERFrequency_top_overlay.png"), 0.72, 1.85, 6.4, 4.25)
    preview(fig("DEERFrequency_top_overlay.png"))
    deerf = rows_for(best, "DEERFrequency", 5)
    add_table(
        slide,
        ["Rank", "Center", "FWHM", "R2"],
        [[r["rank"], fnum(r["center_freq_units"]), fnum(r["fwhm_units"]), fnum(r["r2"])] for r in deerf],
        7.5,
        2.05,
        4.75,
        1.45,
        8.3,
    )
    add_bullets(
        slide,
        [
            "Cleaner sweeps center at ~0.7414-0.7417 in scan units.",
            "Multi-peak fits flag broader or more structured scans; they are useful for exploration but less direct for a single resonance claim.",
            "Frequency-sweep R2 is moderate, so resonance assignment should be paired with duration/decay controls.",
        ],
        7.55,
        3.92,
        4.55,
        1.65,
        13.2,
    )
    footer(slide)

    # 16 DEER duration
    slide = new_slide("DEER duration sweeps")
    title(slide, "deer", "DEER pulse-duration sweeps show usable but non-ideal RF2 control")
    image_fit(slide, fig("DEERDuration_top_overlay.png"), 0.72, 1.85, 6.45, 4.25)
    preview(fig("DEERDuration_top_overlay.png"))
    deerd = rows_for(best, "DEERDuration", 5)
    add_bar_chart(slide, "DEER RF2 pi time [ns]", [str(i) for i in range(1, len(deerd) + 1)], [float(r["pi_time_ns"]) for r in deerd], 7.55, 2.0, 4.75, 2.55, ACCENT2)
    add_text(slide, "Best duration sweeps suggest RF2 pi times near 38-47 ns, but median R2 is only ~0.51. Treat these as calibration guides, then verify with on/off-resonance DEER decay contrast.", 7.55, 4.88, 4.65, 0.95, 13.2, INK)
    footer(slide)

    # 17 DEER decay
    slide = new_slide("DEER decay")
    title(slide, "deer", "DEER decay scans are the cleanest DEER evidence layer")
    image_fit(slide, fig("DEERDecay_top_overlay.png"), 0.72, 1.85, 6.45, 4.25)
    preview(fig("DEERDecay_top_overlay.png"))
    deerc = rows_for(best, "DEERDecay", 5)
    add_bar_chart(slide, "DEER decay T2-like time [ns]", [str(i) for i in range(1, len(deerc) + 1)], [float(r["T2_ns"]) for r in deerc], 7.55, 2.0, 4.75, 2.55, BLUE)
    add_bullets(
        slide,
        [
            "DEER decay fits are high quality: median R2 ~0.985.",
            "On/off-resonance controls sit in the ~625-800 ns range, making them credible for comparative sensing.",
            "The May13 on-resonance 200 ns scan is the best single DEER decay evidence object.",
        ],
        7.55,
        4.88,
        4.65,
        1.45,
        13.2,
    )
    footer(slide)

    # 18 Iteration stability
    slide = new_slide("Iteration stability")
    title(slide, "stability", "Iteration-level data separates real physics from acquisition drift")
    ranked_iter = sorted(iteration, key=lambda r: abs(float(r["first_last_delta"] or 0)), reverse=True)[:7]
    add_table(
        slide,
        ["Experiment", "Iterations", "mean", "delta"],
        [[r["experiment_type"], r["n_iterations"], fnum(r["mean_contrast"]), fnum(r["first_last_delta"])] for r in ranked_iter],
        0.78,
        1.95,
        5.05,
        2.2,
        8.1,
    )
    first_fig = ANALYSIS / ranked_iter[0]["figure"].replace("/", "\\") if ranked_iter and ranked_iter[0].get("figure") else Path("")
    image_fit(slide, first_fig, 6.2, 1.85, 6.05, 4.25)
    preview(first_fig)
    add_text(slide, "Use these plots to reject runs where the average trace is shaped by drift or intermittent refocusing/optimization rather than stable sequence response.", 0.86, 4.72, 4.8, 0.95, 14, INK)
    footer(slide)

    # 19 Spatial context
    slide = new_slide("Spatial and alignment context")
    title(slide, "context", "Spatial scans explain why some sequence families move in and out of the good regime")
    scans = [r for r in inventory if r["experiment_type"] == "LineScan"][:6]
    for idx, r in enumerate(scans[:4]):
        x = 0.75 + (idx % 2) * 5.95
        y = 1.85 + (idx // 2) * 2.25
        image_fit(slide, ANALYSIS / r["figure"].replace("/", "\\"), x, y, 5.45, 1.9)
    if scans:
        preview(ANALYSIS / scans[0]["figure"].replace("/", "\\"))
    add_text(slide, "Alignment data is not the main scientific result, but it is essential context for interpreting day-to-day contrast, collection, and reference-peak changes.", 0.85, 6.58, 11.4, 0.35, 12.5, INK)
    footer(slide)

    # 20 Cross-sequence comparison
    slide = new_slide("Cross-sequence time scales")
    title(slide, "synthesis", "The hierarchy is short Ramsey, microsecond echo/DD, and much longer T1")
    cats = ["Ramsey T2*", "Spin echo T2", "XY8/DD T2", "DEER decay", "T1 best"]
    vals_us = [
        float(ramsey[0]["T2_star_ns"]) / 1000.0,
        float(echo[0]["T2_ns"]) / 1000.0,
        float(dd[0]["T2_dd_ns"]) / 1000.0,
        float(deerc[0]["T2_ns"]) / 1000.0,
        float(t1[0]["T1_ns"]) / 1000.0,
    ]
    add_bar_chart(slide, "Best extracted characteristic time [us]", cats, vals_us, 0.75, 1.9, 7.4, 4.25, ACCENT)
    add_bullets(
        slide,
        [
            "Free precession is the limiting regime: T2* ~0.17 us in the best Ramsey run.",
            "Refocusing moves useful coherence to ~0.8-1.2 us; XY8 can extend the envelope but needs careful confidence checks.",
            "Relaxation is much longer when measured with the credible normalized/balanced T1 protocol.",
        ],
        8.55,
        2.05,
        3.7,
        2.4,
        15,
    )
    footer(slide)

    # 21 Best-run recommendations
    slide = new_slide("Recommended reference runs")
    title(slide, "action", "Use a small set of high-value files as the working reference package")
    recs = [
        ["ODMR", "May2 1750 ZFS", "center ~2.868, strong fit"],
        ["Rabi", "May11 1205 after combiner", "pi ~24.6 ns, R2 ~0.987"],
        ["Ramsey", "May9 0125 10 MHz detune", "T2* ~174 ns"],
        ["Spin echo", "May11 0023 / May13 0436", "T2 ~0.78-0.87 us"],
        ["XY8", "May11 XY8-1/4/8", "DD scaling reference"],
        ["DEER", "May13 decay + May11/12 freq", "resonance and decay controls"],
        ["T1", "May4 NICR balanced", "T1 ~254 us"],
    ]
    add_table(slide, ["Family", "Reference run", "Why it matters"], recs, 0.8, 1.9, 11.75, 3.9, 10.5)
    add_text(slide, "These are not the only useful files; they are the first-pass set a student should reopen interactively before collecting the next experimental block.", 0.9, 6.05, 10.7, 0.5, 14, INK)
    footer(slide)

    # 22 Interpretation caveats
    slide = new_slide("Caveats and scientific risks")
    title(slide, "caveats", "The analysis is broad, but not every numerical fit should be promoted to a physical claim")
    add_bullets(
        slide,
        [
            "Repeated ODMR alignment scans are context; they should not inflate confidence in resonance assignment.",
            "Checkpoint files can be scientifically useful but need completion metadata reviewed before final publication use.",
            "Low-R2 Ramsey and DEER duration/frequency sweeps may encode real behavior plus acquisition artifacts; treat them as qualitative unless reproduced.",
            "T1 is protocol-sensitive in this dataset; balanced/normalized runs are credible, raw long scans are not.",
            "Single-point XY8 noise-spectrum files are response samples, not decay curves.",
        ],
        0.9,
        1.95,
        5.6,
        4.0,
        18,
    )
    image_fit(slide, fig("nv_metric_distributions.png"), 6.95, 1.85, 5.45, 4.15)
    preview(fig("nv_metric_distributions.png"))
    footer(slide)

    # 23 Next experiments
    slide = new_slide("Next experiments")
    title(slide, "next", "The next data block should reduce ambiguity, not just add more files")
    add_table(
        slide,
        ["Need", "Experiment", "Decision it unlocks"],
        [
            ["Pulse calibration", "Repeat Rabi at final RF power and stage position", "Lock pi/pi-half for all sensing sequences"],
            ["Dephasing", "Ramsey +/-X with explicit detuning ladder", "Separate phase control from T2* extraction"],
            ["Coherence", "Spin echo + XY8-N repeated in one session", "True DD scaling without day-to-day drift"],
            ["DEER", "Freq sweep -> duration sweep -> on/off decay in one block", "Validate RF2 resonance and coupling contrast"],
            ["T1", "Balanced normalized T1 repeat", "Constrain relaxation with credible model family"],
        ],
        0.78,
        1.9,
        11.75,
        3.55,
        10.3,
    )
    add_text(slide, "Practical principle: fewer, chained, internally controlled runs will be more valuable than another large folder of independent exploratory files.", 0.9, 5.9, 10.8, 0.45, 15, ACCENT, True)
    footer(slide)

    # 24 Appendix index
    slide = new_slide("Appendix and file index")
    title(slide, "appendix", "The full analysis package remains the source of detailed evidence")
    add_bullets(
        slide,
        [
            "HTML report: report_429_TDD.html",
            "Fit table: tables/fit_summary.csv",
            "Best metrics: tables/best_metrics.csv",
            "Iteration stability: tables/iteration_summary.csv",
            "Noise-spectrum points: tables/noise_spectrum_summary.csv",
            "Per-file plots: figures/per_file/",
            "Curated overlays: figures/curated/",
            "Diagnostics: figures/diagnostics/",
        ],
        0.95,
        1.95,
        5.2,
        3.9,
        17,
    )
    image_fit(slide, fig("DEERDecay_top_overlay.png"), 6.75, 1.85, 5.65, 4.25)
    preview(fig("DEERDecay_top_overlay.png"))
    footer(slide)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(PPTX_PATH)
    return slide_titles, preview_imgs


def main() -> None:
    if OUT_DIR.exists():
        OUT_DIR.mkdir(parents=True, exist_ok=True)
    slide_titles, preview_imgs = build_deck()
    contact = make_preview_contact(slide_titles, preview_imgs)
    build_info = {
        "pptx": str(PPTX_PATH),
        "slide_count": len(slide_titles),
        "preview_contact_sheet": str(contact),
        "source_analysis": str(ANALYSIS),
    }
    (OUT_DIR / "deck_build_info.json").write_text(json.dumps(build_info, indent=2), encoding="utf-8")
    print(json.dumps(build_info, indent=2))


if __name__ == "__main__":
    main()
